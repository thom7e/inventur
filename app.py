import io
import os
import re
import secrets
from datetime import datetime, timedelta, timezone
from functools import wraps
import sqlite3
import logging
from contextlib import contextmanager
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from uuid import uuid4

import openpyxl
from flask import (
    Flask,
    Response,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.security import check_password_hash, generate_password_hash


def _generate_default_secret() -> str:
    return os.urandom(32).hex()


def quantize_currency(value: Decimal) -> Decimal:
    return value.quantize(CURRENCY_QUANTIZER, rounding=ROUND_HALF_UP)


def quantize_quantity(value: Decimal) -> Decimal:
    return value.quantize(QUANTITY_QUANTIZER, rounding=ROUND_HALF_UP)


def format_quantity_for_display(value: Decimal) -> str:
    formatted = f"{value:.3f}"
    if "." in formatted:
        formatted = formatted.rstrip("0").rstrip(".")
    return formatted or "0"


def format_currency_for_display(value: Decimal) -> str:
    return f"{quantize_currency(value):.2f}"


def normalize_decimal_value(value, *, default=Decimal('0')) -> Decimal:
    if value is None:
        return default
    normalized = str(value).strip()
    if not normalized:
        return default
    normalized = normalized.replace(',', '.')
    try:
        return Decimal(normalized)
    except InvalidOperation:
        filtered = ''.join(ch for ch in normalized if ch.isdigit() or ch in '.-')
        if not filtered or filtered in {'.', '-', '-.'}:
            return default
        try:
            return Decimal(filtered)
        except InvalidOperation:
            return default


def respond_error(message: str, status: int = 400):
    return jsonify({"success": False, "message": message}), status


logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# CSRF-Schutz
# ---------------------------------------------------------------------------

def generate_csrf_token() -> str:
    """Gibt den Session-CSRF-Token zurück (wird bei Bedarf erstellt)."""
    if 'csrf_token' not in session:
        session['csrf_token'] = secrets.token_hex(32)
    return session['csrf_token']


def validate_csrf() -> bool:
    """Prüft den CSRF-Token aus Formular, JSON-Body oder HTTP-Header."""
    expected = session.get('csrf_token')
    if not expected:
        return False
    token = (
        request.form.get('_csrf_token')
        or request.headers.get('X-CSRF-Token')
        or (request.get_json(silent=True) or {}).get('_csrf_token')
    )
    return secrets.compare_digest(str(token or ''), str(expected))


def require_csrf(f):
    """Dekorator: Validiert den CSRF-Token bei zustandsändernden Methoden."""
    @wraps(f)
    def decorated(*args, **kwargs):
        if request.method in ('POST', 'PUT', 'PATCH', 'DELETE'):
            if not validate_csrf():
                if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return respond_error('Ungültiger CSRF-Token.', status=403)
                return Response('Ungültiger CSRF-Token.', status=403)
        return f(*args, **kwargs)
    return decorated


def require_admin(f):
    """Dekorator: Erfordert eine aktive Admin-Session."""
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('is_admin'):
            if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return respond_error('Unbefugter Zugriff', status=403)
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated


# ---------------------------------------------------------------------------
# Audit-Log
# ---------------------------------------------------------------------------

def ensure_audit_log_table() -> None:
    with db_connection(COST_CENTER_DB) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS audit_log (
                id        INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT    NOT NULL,
                user      TEXT    NOT NULL,
                action    TEXT    NOT NULL,
                details   TEXT
            )
        """)
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_audit_log_ts ON audit_log(timestamp)"
        )
        conn.commit()


def audit_log(action: str, details: str = '') -> None:
    """Schreibt einen Audit-Log-Eintrag in die Datenbank und den Logger."""
    user = 'admin' if session.get('is_admin') else 'anonym'
    ts = datetime.now(timezone.utc).isoformat()
    logger.info("AUDIT | %s | %s | %s | %s", ts, user, action, details)
    try:
        with db_connection(COST_CENTER_DB) as conn:
            conn.execute(
                "INSERT INTO audit_log (timestamp, user, action, details) VALUES (?, ?, ?, ?)",
                (ts, user, action, details),
            )
            conn.commit()
    except sqlite3.Error:
        logger.exception("Audit-Log konnte nicht geschrieben werden")


@contextmanager
def db_connection(db_path: str):
    conn = sqlite3.connect(db_path, timeout=10)
    conn.row_factory = sqlite3.Row  # wichtig: Zugriff per row["feld"]
    try:
        yield conn
    finally:
        conn.close()


def ensure_cart_table() -> None:
    with db_connection(COST_CENTER_DB) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS session_cart_items (
                cart_token TEXT NOT NULL,
                artikelname TEXT NOT NULL,
                menge TEXT NOT NULL,
                einheit TEXT,
                preis TEXT NOT NULL,
                gesamtpreis TEXT NOT NULL,
                PRIMARY KEY (cart_token, artikelname)
            )
            """
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_session_cart_token ON session_cart_items(cart_token)"
        )
        conn.commit()


def ensure_cost_center_archive_table() -> None:
    with db_connection(COST_CENTER_DB) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS kostenstellen_archive (
                archiv_jahr INTEGER NOT NULL,
                kostenstellen_nummer TEXT NOT NULL,
                kostenstellen_bezeichnung TEXT NOT NULL,
                archived_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_kostenstellen_archive_jahr ON kostenstellen_archive(archiv_jahr)")

        # Archiv schreibschützen
        conn.execute("""
            CREATE TRIGGER IF NOT EXISTS ro_kst_arch_ins
            BEFORE INSERT ON kostenstellen_archive
            BEGIN
                SELECT RAISE(ABORT, 'kostenstellen_archive ist schreibgeschützt');
            END;
        """)
        conn.execute("""
            CREATE TRIGGER IF NOT EXISTS ro_kst_arch_upd
            BEFORE UPDATE ON kostenstellen_archive
            BEGIN
                SELECT RAISE(ABORT, 'kostenstellen_archive ist schreibgeschützt');
            END;
        """)
        conn.execute("""
            CREATE TRIGGER IF NOT EXISTS ro_kst_arch_del
            BEFORE DELETE ON kostenstellen_archive
            BEGIN
                SELECT RAISE(ABORT, 'kostenstellen_archive ist schreibgeschützt');
            END;
        """)
        conn.commit()

def ensure_app_meta_table() -> None:
    with db_connection(COST_CENTER_DB) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS app_meta (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
            """
        )
        conn.commit()

def ensure_cost_center_unique_index():
    with db_connection(COST_CENTER_DB) as conn:
        conn.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS ux_kostenstellen_nummer
            ON kostenstellen (kostenstellen_nummer COLLATE NOCASE)
        """)
        conn.commit()

def get_active_inventur_year() -> int:
    # 1) aus app_meta lesen
    with db_connection(COST_CENTER_DB) as conn:
        row = conn.execute("SELECT value FROM app_meta WHERE key = 'active_inventur_year'").fetchone()
        if row and str(row["value"]).strip():
            try:
                return int(str(row["value"]).strip())
            except Exception:
                pass

        # 2) Fallback: benutze das bisher höchste Jahr aus Inventur (falls vorhanden)
        row2 = conn.execute("SELECT MAX(jahr) AS max_jahr FROM inventur").fetchone()
        if row2 and row2["max_jahr"] is not None:
            try:
                return int(row2["max_jahr"])
            except Exception:
                pass

    # 3) letzter Fallback: Kalenderjahr
    return datetime.now().year


def set_active_inventur_year(year: int) -> None:
    with db_connection(COST_CENTER_DB) as conn:
        conn.execute(
            "INSERT INTO app_meta(key, value) VALUES('active_inventur_year', ?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (str(int(year)),)
        )
        conn.commit()



def parse_decimal(
    value,
    field_name: str,
    *,
    allow_negative: bool = False,
    allow_zero: bool = False,
) -> Decimal:
    if value is None:
        raise ValueError(f"{field_name} ist erforderlich.")
    normalized = str(value).replace(",", ".").strip()
    if not normalized:
        raise ValueError(f"{field_name} ist erforderlich.")
    try:
        parsed = Decimal(normalized)
    except InvalidOperation as exc:
        raise ValueError(f"{field_name} hat ein ungültiges Format.") from exc
    if parsed == 0 and not allow_zero:
        raise ValueError(f"{field_name} darf nicht 0 sein.")
    if parsed < 0 and not allow_negative:
        raise ValueError(f"{field_name} darf nicht negativ sein.")
    return parsed


def fetch_article_by_name(artikelname: str):
    with db_connection(ARTICLE_DB) as conn:
        row = conn.execute(
            "SELECT artikelname, einheit, preis FROM artikel WHERE artikelname = ? COLLATE NOCASE",
            (artikelname.strip(),),
        ).fetchone()
    return row


def fetch_articles(search_query: str):
    term = f"%{search_query.strip().lower()}%"
    with db_connection(ARTICLE_DB) as conn:
        cur = conn.execute(
            "SELECT artikelname, einheit, preis FROM artikel WHERE lower(artikelname) LIKE ? ORDER BY artikelname LIMIT 500",
            (term,),
        )
        return cur.fetchall()


def serialize_article_row(row) -> dict:
    # sicherer Feldzugriff fuer sqlite3.Row / dict / obj
    def _get(r, name):
        try:
            return r[name]
        except Exception:
            try:
                return getattr(r, name)
            except Exception:
                try:
                    return r.get(name) if hasattr(r, "get") else None
                except Exception:
                    return None

    preis_decimal = quantize_currency(normalize_decimal_value(_get(row, "preis")))

    artikelname = _get(row, "artikelname") or ""
    einheit = _get(row, "einheit") or ""

    return {
        "artikelname": artikelname,
        "einheit": einheit,
        "preis": float(preis_decimal),
    }


def article_rows_to_payload(rows):
    return [serialize_article_row(row) for row in rows]


def fetch_kostenstellen():
    with db_connection(COST_CENTER_DB) as conn:
        rows = conn.execute(
            "SELECT kostenstellen_nummer, kostenstellen_bezeichnung FROM kostenstellen ORDER BY kostenstellen_nummer"
        ).fetchall()
    return rows


def fetch_cart_items(cart_token: str):
    if not cart_token:
        return []
    with db_connection(COST_CENTER_DB) as conn:
        rows = conn.execute(
            """
            SELECT artikelname, menge, einheit, preis, gesamtpreis
            FROM session_cart_items
            WHERE cart_token = ?
            ORDER BY LOWER(artikelname)
            """,
            (cart_token,),
        ).fetchall()
    cart_items = []
    for row in rows:
        cart_items.append(
            {
                "artikelname": row["artikelname"],
                "menge": Decimal(row["menge"]),
                "einheit": row["einheit"],
                "preis": Decimal(row["preis"]),
                "gesamtpreis": Decimal(row["gesamtpreis"]),
            }
        )
    return cart_items


def cart_items_for_template(cart_items):
    return [
        {
            "artikelname": item["artikelname"],
            "menge": format_quantity_for_display(item["menge"]),
            "einheit": item["einheit"],
            "preis": format_currency_for_display(item["preis"]),
            "gesamtpreis": format_currency_for_display(item["gesamtpreis"]),
        }
        for item in cart_items
    ]


def cart_items_for_json(cart_items):
    return cart_items_for_template(cart_items)


def update_cart_item(
    cart_token: str,
    artikelname: str,
    einheit: str,
    unit_price: Decimal,
    menge_delta: Decimal,
) -> None:
    unit_price = quantize_currency(unit_price)
    menge_delta = quantize_quantity(menge_delta)
    if menge_delta == 0:
        return
    with db_connection(COST_CENTER_DB) as conn:
        row = conn.execute(
            "SELECT menge FROM session_cart_items WHERE cart_token = ? AND artikelname = ?",
            (cart_token, artikelname),
        ).fetchone()
        if row:
            current_qty = Decimal(row["menge"])
            new_qty = quantize_quantity(current_qty + menge_delta)
            if new_qty <= 0:
                conn.execute(
                    "DELETE FROM session_cart_items WHERE cart_token = ? AND artikelname = ?",
                    (cart_token, artikelname),
                )
            else:
                line_total = quantize_currency(new_qty * unit_price)
                conn.execute(
                    """
                    UPDATE session_cart_items
                    SET menge = ?, einheit = ?, preis = ?, gesamtpreis = ?
                    WHERE cart_token = ? AND artikelname = ?
                    """,
                    (
                        str(new_qty),
                        einheit,
                        str(unit_price),
                        str(line_total),
                        cart_token,
                        artikelname,
                    ),
                )
        else:
            if menge_delta <= 0:
                conn.commit()
                return
            line_total = quantize_currency(menge_delta * unit_price)
            conn.execute(
                """
                INSERT INTO session_cart_items (cart_token, artikelname, menge, einheit, preis, gesamtpreis)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    cart_token,
                    artikelname,
                    str(menge_delta),
                    einheit,
                    str(unit_price),
                    str(line_total),
                ),
            )
        conn.commit()


def clear_cart_items(cart_token: str) -> None:
    with db_connection(COST_CENTER_DB) as conn:
        conn.execute("DELETE FROM session_cart_items WHERE cart_token = ?", (cart_token,))
        conn.commit()


def group_inventur_rows(rows):
    grouped = {}
    for row in rows:
        # sqlite3.Row -> dict (inkl. _rid, falls selektiert)
        if hasattr(row, "keys"):
            row_dict = {k: row[k] for k in row.keys()}
        else:
            # Fallback: per Index (nur falls wirklich kein Row-Objekt)
            try:
                row_dict = {
                    "_rid":        row[0] if len(row) > 0 else None,
                    "kostenstelle": row[1],
                    "jahr":         row[2],
                    "artikelname":  row[3],
                    "einheit":      row[4],
                    "menge":        row[5],
                    "preis":        row[6],
                    "gesamtpreis":  row[7],
                }
            except Exception:
                continue

        ks = (row_dict.get("kostenstelle") or "").strip()
        entry = grouped.setdefault(ks, {"materialien": [], "gesamtsumme": Decimal("0")})
        entry["materialien"].append(row_dict)

        total_decimal = normalize_decimal_value(row_dict.get("gesamtpreis"), default=Decimal("0"))
        entry["gesamtsumme"] += quantize_currency(total_decimal)

    return grouped


app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.getenv("FLASK_SECRET_KEY") or _generate_default_secret(),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
)

ARTICLE_DB = os.getenv("ARTICLE_DB_PATH", "artikel.db")
COST_CENTER_DB = os.getenv("COST_CENTER_DB_PATH", "kostenstellen_db.db")
CURRENCY_QUANTIZER = Decimal("0.01")
QUANTITY_QUANTIZER = Decimal("0.001")

ADMIN_USERNAME = os.getenv("ADMIN_USERNAME", "admin")
_admin_password_hash = os.getenv("ADMIN_PASSWORD_HASH")
if not _admin_password_hash:
    _admin_password_hash = generate_password_hash(os.getenv("ADMIN_PASSWORD", "admin"))
ADMIN_PASSWORD_HASH = _admin_password_hash


def _bootstrap() -> None:
    try:
        ensure_cart_table()
        ensure_cost_center_archive_table()
        ensure_app_meta_table()
        ensure_cost_center_unique_index()
        ensure_audit_log_table()
    except Exception:
        app.logger.exception("Bootstrap failed")


# Register bootstrap for Flask <3 and call immediately for Flask 3+
if hasattr(app, "before_first_request"):
    app.before_first_request(_bootstrap)  # type: ignore[attr-defined]
else:
    _bootstrap()


@app.context_processor
def inject_csrf_token():
    """Macht csrf_token() in allen Templates verfügbar."""
    return {'csrf_token': generate_csrf_token}


@app.before_request
def persist_session() -> None:
    session.permanent = True
    if "cart_token" not in session:
        session["cart_token"] = uuid4().hex


@app.route('/login', methods=['GET', 'POST'])
@require_csrf
def login():
    error = None
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        if username != ADMIN_USERNAME or not check_password_hash(ADMIN_PASSWORD_HASH, password):
            error = 'Ungültiger Benutzername oder Passwort'
            logger.warning("Fehlgeschlagener Login-Versuch für Benutzer: %s", username)
        else:
            audit_log('admin_login', f'Benutzer: {username}')
            session.clear()
            session['is_admin'] = True
            session['cart_token'] = uuid4().hex
            return redirect(url_for('admin'))
    return render_template('login.html', error=error)


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


@app.route('/admin')
@require_admin
def admin():

    with db_connection(ARTICLE_DB) as conn:
        rows = conn.execute(
            'SELECT artikelname, einheit, preis FROM artikel ORDER BY artikelname'
        ).fetchall()
    articles = article_rows_to_payload(rows)

    # NEU: Jahre bestimmen
    active_year = get_active_inventur_year()
    existing_years = fetch_available_years()           # Jahre, die es in inventur gibt
    archived_years = set(fetch_archived_years())       # bereits archiviert

    # Jahre, die existieren und noch nicht archiviert sind
    closable_years = sorted([y for y in set(existing_years or [active_year]) if y not in archived_years], reverse=True)

    # Fallback: wenn gar nichts drin ist, mind. aktives Jahr anbieten
    if not closable_years:
        closable_years = [active_year]

    return render_template(
        'admin.html',
        articles=articles,
        current_year=datetime.now().year,
        active_inventur_year=active_year,   # <-- neu
        closable_years=closable_years       # <-- neu
    )

@app.route('/admin/clear_cost_centers', methods=['POST'])
@require_admin
@require_csrf
def admin_clear_cost_centers():
    with db_connection(COST_CENTER_DB) as conn:
        cur = conn.execute('DELETE FROM kostenstellen')
        deleted = cur.rowcount if hasattr(cur, 'rowcount') else None
        conn.commit()

    msg = 'Kostenstellen wurden geleert.'
    if isinstance(deleted, int) and deleted >= 0:
        msg = f'Kostenstellen wurden geleert ({deleted} Einträge).'
    audit_log('kostenstellen_geleert', f'{deleted if deleted is not None else "?"} Einträge gelöscht')
    return jsonify({'success': True, 'message': msg})


@app.route('/admin/add_cost_center', methods=['POST'])
@require_admin
@require_csrf
def admin_add_cost_center():
    data = request.get_json(silent=True) or {}
    nummer = (data.get('nummer') or '').strip()
    bezeichnung = (data.get('bezeichnung') or '').strip()

    if not nummer or not bezeichnung:
        return respond_error('Nummer und Bezeichnung sind erforderlich.')

    nummer_norm = nummer

    with db_connection(COST_CENTER_DB) as conn:
        exists = conn.execute(
            'SELECT 1 FROM kostenstellen WHERE kostenstellen_nummer = ? COLLATE NOCASE',
            (nummer_norm,)
        ).fetchone()
        if exists:
            return respond_error('Diese Kostenstellen-Nummer existiert bereits.')

        conn.execute(
            'INSERT INTO kostenstellen (kostenstellen_nummer, kostenstellen_bezeichnung) VALUES (?, ?)',
            (nummer_norm, bezeichnung)
        )
        conn.commit()

    audit_log('kostenstelle_hinzugefuegt', f'Nummer: {nummer_norm}, Bezeichnung: {bezeichnung}')
    return jsonify({'success': True, 'message': f'Kostenstelle {nummer_norm} hinzugefügt.'})

@app.route('/admin/bulk_add_cost_centers', methods=['POST'])
@require_admin
@require_csrf
def admin_bulk_add_cost_centers():

    data = request.get_json(silent=True) or {}
    entries = data.get('kostenstellen', [])
    if not isinstance(entries, list) or not entries:
        return respond_error('Keine Kostenstellen übermittelt.')

    # Payload validieren & in (nummer, bezeichnung) normalisieren
    cleaned = []
    seen = set()
    for e in entries:
        nummer = (e.get('nummer') or '').strip()
        bezeichnung = (e.get('bezeichnung') or '').strip()
        if not nummer or not bezeichnung:
            return respond_error('Jede Zeile benötigt "Nummer" und "Bezeichnung".')
        key = nummer.lower()
        if key in seen:
            # Duplikate innerhalb der Liste ignorieren
            continue
        seen.add(key)
        cleaned.append((nummer, bezeichnung))

    if not cleaned:
        return respond_error('Keine neuen Kostenstellen nach Bereinigung.')

    added, skipped = 0, 0
    with db_connection(COST_CENTER_DB) as conn:
        for nummer, bezeichnung in cleaned:
            # existiert schon?
            row = conn.execute(
                'SELECT 1 FROM kostenstellen WHERE kostenstellen_nummer = ? COLLATE NOCASE',
                (nummer,)
            ).fetchone()
            if row:
                skipped += 1
                continue
            conn.execute(
                'INSERT INTO kostenstellen (kostenstellen_nummer, kostenstellen_bezeichnung) VALUES (?, ?)',
                (nummer, bezeichnung)
            )
            added += 1
        conn.commit()

    msg = f'{added} Kostenstellen hinzugefügt'
    if skipped:
        msg += f', {skipped} übersprungen (bereits vorhanden)'
    audit_log('kostenstellen_bulk_hinzugefuegt', f'{added} hinzugefügt, {skipped} übersprungen')
    return jsonify({'success': True, 'message': msg, 'added': added, 'skipped': skipped})


@app.route('/admin/inventur/update_row', methods=['POST'])
@require_admin
@require_csrf
def admin_update_inventur_row():

    data = request.get_json(silent=True) or {}
    try:
        rid = int(data.get('id'))
    except (TypeError, ValueError):
        return respond_error('Ungültige ID.')

    new_name  = (data.get('artikelname') or '').strip()
    new_unit  = (data.get('einheit') or '').strip()
    menge_raw = data.get('menge')
    preis_raw = data.get('preis')
    if not new_name or not new_unit:
        return respond_error('Artikelname und Einheit sind erforderlich.')
    try:
        menge = parse_decimal(menge_raw, 'Menge', allow_zero=False)
        preis = parse_decimal(preis_raw, 'Preis', allow_zero=True)
    except ValueError as exc:
        return respond_error(str(exc))

    active_year = get_active_inventur_year()
    with db_connection(COST_CENTER_DB) as conn:
        row = conn.execute("SELECT jahr FROM inventur WHERE rowid = ?", (rid,)).fetchone()
        if not row:
            return respond_error('Eintrag nicht gefunden.', status=404)
        if int(row['jahr']) != int(active_year):
            return respond_error('Bearbeiten ist nur im aktiven Jahr erlaubt.', status=403)

        gesamt = quantize_currency(menge * quantize_currency(preis))
        conn.execute(
            """UPDATE inventur
               SET artikelname = ?, einheit = ?, menge = ?, preis = ?, gesamtpreis = ?
             WHERE rowid = ?""",
            (new_name, new_unit, float(quantize_quantity(menge)),
             float(quantize_currency(preis)), float(gesamt), rid)
        )
        conn.commit()

    audit_log('inventur_zeile_aktualisiert', f'ID: {rid}, Artikel: {new_name}')
    return jsonify({'success': True, 'message': 'Eintrag aktualisiert.'})

@app.route('/admin/inventur/delete_row', methods=['POST'])
@require_admin
@require_csrf
def admin_delete_inventur_row():
    data = request.get_json(silent=True) or {}
    try:
        rid = int(data.get('id'))
    except (TypeError, ValueError):
        return respond_error('Ungültige ID.')

    active_year = get_active_inventur_year()
    with db_connection(COST_CENTER_DB) as conn:
        row = conn.execute("SELECT jahr FROM inventur WHERE rowid = ?", (rid,)).fetchone()
        if not row:
            return respond_error('Eintrag nicht gefunden.', status=404)
        if int(row['jahr']) != int(active_year):
            return respond_error('Löschen ist nur im aktiven Jahr erlaubt.', status=403)

        conn.execute("DELETE FROM inventur WHERE rowid = ?", (rid,))
        conn.commit()

    audit_log('inventur_zeile_geloescht', f'ID: {rid}')
    return jsonify({'success': True, 'message': 'Eintrag gelöscht.'})



@app.route('/get_cart')
def get_cart():
    cart_items = fetch_cart_items(session.get('cart_token'))
    return jsonify(cart_items_for_json(cart_items))


@app.route('/', methods=['GET', 'POST'])
def index():
    cart_token = session.get('cart_token')
    cart_items = fetch_cart_items(cart_token)
    cart_template_items = cart_items_for_template(cart_items)

    # NUR aktive Kostenstellen aus der akt. Tabelle (nicht Archiv)
    rows = fetch_kostenstellen()
    kostenstellen = [(r["kostenstellen_nummer"], r["kostenstellen_bezeichnung"]) for r in rows]

    selected_cost_center = request.form.get('kostenstelle') if request.method == 'POST' else None
    search_query = request.form.get('search_query', '').strip() if request.method == 'POST' else ''

    # >>> diese beiden IMMER definieren
    results_payload = []
    cart_payload = cart_items_for_json(cart_items)

    if request.method == 'POST' and search_query:
        results_rows = fetch_articles(search_query)
        results_payload = article_rows_to_payload(results_rows)
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify(results_payload)

    return render_template(
        'index.html',
        results_payload=results_payload,
        cart_payload=cart_payload,
        search_query=search_query,
        kostenstellen=kostenstellen,
        ausgewaehlte_kostenstelle=selected_cost_center,
        cart=cart_template_items,
        active_inventur_year=get_active_inventur_year(),
    )


@app.route('/search', methods=['POST'])
def search():
    # akzeptiere JSON oder Form-POST
    data = request.get_json(silent=True) or request.form or {}
    query = (data.get("search_query") if isinstance(data, dict) else "") or ""
    query = str(query).strip()
    if len(query) < 2:
        return jsonify({"results": [], "count": 0})
    try:
        rows = fetch_articles(query)
        payload = article_rows_to_payload(rows)
        return jsonify({"results": payload, "count": len(payload)})
    except Exception:
        logger.exception("Fehler bei Suche")
        return jsonify({"results": [], "count": 0, "error": "internal"}), 500


@app.route('/admin/search', methods=['POST'])
def admin_search():
    data = request.get_json(silent=True) or request.form or {}
    query = (data.get("search_query") if isinstance(data, dict) else "") or ""
    query = str(query).strip()
    if len(query) < 2:
        return jsonify({"results": [], "count": 0})
    try:
        rows = fetch_articles(query)
        payload = article_rows_to_payload(rows)
        return jsonify({"results": payload, "count": len(payload)})
    except Exception:
        logger.exception("Fehler bei Admin-Suche")
        return jsonify({"results": [], "count": 0, "error": "internal"}), 500


@app.route('/add_to_cart', methods=['POST'])
@require_csrf
def add_to_cart():
    artikelname = request.form.get('artikelname', '').strip()
    menge_raw = request.form.get('menge')

    if not artikelname:
        return respond_error('Artikelname ist erforderlich.')

    try:
        menge = parse_decimal(menge_raw, 'Menge', allow_negative=True)
    except ValueError as exc:
        return respond_error(str(exc))

    article = fetch_article_by_name(artikelname)
    if not article:
        return respond_error('Artikel wurde nicht gefunden.', 404)

    cart_token = session['cart_token']
    unit_price = Decimal(str(article['preis']))
    update_cart_item(cart_token, article['artikelname'], article['einheit'], unit_price, menge)

    cart_items = fetch_cart_items(cart_token)
    return jsonify(cart_items_for_json(cart_items))


@app.route('/add_artikel', methods=['POST'])
@require_csrf
def add_artikel():
    data = request.get_json(silent=True) or {}
    artikelname = (data.get('artikelname') or '').strip()
    einheit = (data.get('einheit') or '').strip()
    preis_raw = data.get('preis')

    if not artikelname or not einheit:
        return respond_error('Artikelname und Einheit sind erforderlich.')

    try:
        preis = parse_decimal(preis_raw, 'Preis')
    except ValueError as exc:
        return respond_error(str(exc))

    with db_connection(ARTICLE_DB) as conn:
        existing = conn.execute(
            'SELECT 1 FROM artikel WHERE artikelname = ? COLLATE NOCASE', (artikelname,)
        ).fetchone()
        if existing:
            return respond_error('Artikelname bereits vorhanden.')

        conn.execute(
            'INSERT INTO artikel (artikelname, einheit, preis) VALUES (?, ?, ?)',
            (artikelname, einheit, float(quantize_currency(preis))),
        )
        conn.commit()

    return jsonify({'success': True, 'message': 'Artikel erfolgreich hinzugefuegt'})


@app.route('/clear_cart', methods=['POST'])
@require_csrf
def clear_cart():
    cart_token = session.get('cart_token')
    clear_cart_items(cart_token)
    session['cart_token'] = uuid4().hex
    return redirect(url_for('index'))


@app.route('/export_excel', methods=['POST'])
@require_csrf
def export_excel():
    kostenstelle_nummer = request.form.get('kostenstelle')

    if not kostenstelle_nummer:
        return respond_error('Kostenstellennummer ist erforderlich.')

    with db_connection(COST_CENTER_DB) as conn:
        kostenstelle_bezeichnung = conn.execute(
            'SELECT kostenstellen_bezeichnung FROM kostenstellen WHERE kostenstellen_nummer = ?',
            (kostenstelle_nummer,),
        ).fetchone()

    if not kostenstelle_bezeichnung:
        return respond_error('Keine gültige Kostenstellenbezeichnung gefunden.', 404)

    cart_items = fetch_cart_items(session.get('cart_token'))
    if not cart_items:
        return respond_error('Warenkorb ist leer.', 400)

    bezeichnung = f"{kostenstelle_nummer} {kostenstelle_bezeichnung['kostenstellen_bezeichnung']}"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = kostenstelle_nummer
    header_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    bold_font = Font(bold=True)

    ws.append(['Artikelname', 'Menge', 'Einheit', 'Preis (EUR)', 'Gesamtpreis (EUR)'])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = bold_font
        cell.alignment = Alignment(horizontal='center')

    for item in cart_items:
        ws.append([
            item['artikelname'],
            float(item['menge']),
            item['einheit'],
            float(quantize_currency(item['preis'])),
            float(quantize_currency(item['gesamtpreis'])),
        ])

    data_end_row = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=data_end_row, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = 'EUR #,##0.00'
            cell.alignment = Alignment(horizontal='center')

    ws.append([])
    ws.append([
        'Gesamtpreis fuer die Baustelle:',
        '',
        '',
        '',
        f"=SUM(E2:E{data_end_row})",
    ])
    for cell in ws[ws.max_row]:
        cell.font = bold_font
        cell.alignment = Alignment(horizontal='center')

    for col_index in range(1, ws.max_column + 1):
        column_letter = get_column_letter(col_index)
        max_length = 0
        for cell in ws[column_letter]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{bezeichnung}.xlsx"'},
    )


@app.route('/save_to_inventur', methods=['POST'])
@require_csrf
def save_to_inventur():
    kostenstelle_nummer = request.form.get('kostenstelle')
    jahr = get_active_inventur_year()


    if not kostenstelle_nummer:
        return respond_error('Kostenstellennummer ist erforderlich.')

    with db_connection(COST_CENTER_DB) as conn:
        kostenstelle_bezeichnung = conn.execute(
            'SELECT kostenstellen_bezeichnung FROM kostenstellen WHERE kostenstellen_nummer = ?',
            (kostenstelle_nummer,),
        ).fetchone()

    if not kostenstelle_bezeichnung:
        return respond_error('Keine gültige Kostenstellenbezeichnung gefunden.', 404)

    cart_items = fetch_cart_items(session.get('cart_token'))
    if not cart_items:
        return respond_error('Warenkorb ist leer.', 400)

    kostenstelle = f"{kostenstelle_nummer}-{kostenstelle_bezeichnung['kostenstellen_bezeichnung']}"

    with db_connection(COST_CENTER_DB) as conn:
        for item in cart_items:
            conn.execute(
                'INSERT INTO inventur (kostenstelle, jahr, artikelname, einheit, menge, preis, gesamtpreis) VALUES (?, ?, ?, ?, ?, ?, ?)',
                (
                    kostenstelle,
                    jahr,
                    item['artikelname'],
                    item['einheit'],
                    float(item['menge']),
                    float(quantize_currency(item['preis'])),
                    float(quantize_currency(item['gesamtpreis'])),
                ),
            )
        conn.commit()

    audit_log('inventur_gespeichert', f'Kostenstelle: {kostenstelle}, Jahr: {jahr}, Artikel: {len(cart_items)}')
    clear_cart_items(session.get('cart_token'))
    session['cart_token'] = uuid4().hex
    return redirect(url_for('index'))


@app.route('/admin/apply_price_factor', methods=['POST'])
@require_admin
@require_csrf
def apply_price_factor():

    data = request.get_json(silent=True) or {}
    factor_raw = data.get('factor')
    rounding_places = data.get('rounding_places', 2)

    try:
        factor = parse_decimal(factor_raw, 'Faktor', allow_negative=False)
    except ValueError as exc:
        return respond_error(str(exc))

    try:
        rounding_places = int(rounding_places)
    except (TypeError, ValueError):
        rounding_places = 2
    rounding_places = max(0, min(4, rounding_places))
    quantizer = Decimal('1').scaleb(-rounding_places)

    with db_connection(ARTICLE_DB) as conn:
        rows = conn.execute('SELECT artikelname, preis FROM artikel').fetchall()
        for row in rows:
            base_price = Decimal(str(row['preis']))
            updated_price = (base_price * factor).quantize(quantizer, rounding=ROUND_HALF_UP)
            conn.execute(
                'UPDATE artikel SET preis = ? WHERE artikelname = ?',
                (float(updated_price), row['artikelname']),
            )
        conn.commit()

    audit_log('preisfaktor_angewendet', f'Faktor: {factor}, Nachkommastellen: {rounding_places}')
    return jsonify({'success': True, 'message': 'Preise erfolgreich aktualisiert'})


@app.route('/admin/year_end', methods=['POST'])
@require_admin
@require_csrf
def perform_year_end():

    data = request.get_json(silent=True) or {}
    jahr_raw = data.get('jahr')
    kostenstellen = data.get('kostenstellen', [])

    try:
        jahr = int(str(jahr_raw).strip())
    except (TypeError, ValueError):
        return respond_error('Bitte eine gültige Jahreszahl angeben.')

    active_year = get_active_inventur_year()
    existing_years = set(fetch_available_years())
    archived_years = set(fetch_archived_years())

    # nur Jahre, die es gibt und noch nicht archiviert sind
    if jahr not in existing_years or jahr in archived_years:
        return respond_error('Dieses Jahr kann nicht abgeschlossen werden (existiert nicht oder ist bereits archiviert).')

    # neue Kostenstellen NUR fürs aktive Jahr zulassen
    if jahr != active_year:
        if kostenstellen:
            return respond_error('Neue Kostenstellen dürfen nur für das aktive Jahr angegeben werden.')
        # für Nicht-Aktivjahr: leere Liste setzen, damit unten keine Fehler kommen
        kostenstellen = []

    # ab hier: bisherige Logik, aber "kostenstellen" kann leer sein
    if jahr == active_year:
        if not isinstance(kostenstellen, list) or not kostenstellen:
            return respond_error('Neue Kostenstellen werden benötigt (für das aktive Jahr).')

        parsed_entries = []
        seen_numbers = set()
        for entry in kostenstellen:
            nummer = (entry.get('nummer') or '').strip()
            bezeichnung = (entry.get('bezeichnung') or '').strip()
            if not nummer or not bezeichnung:
                return respond_error('Jede Kostenstelle benötigt Nummer und Bezeichnung.')
            if nummer in seen_numbers:
                return respond_error(f'Kostenstelle {nummer} ist mehrfach vorhanden.')
            seen_numbers.add(nummer)
            parsed_entries.append((nummer, bezeichnung))
    else:
        parsed_entries = []  # keine neuen KSt. für Nicht-Aktivjahr

    ensure_cost_center_archive_table()

    
    with db_connection(COST_CENTER_DB) as conn:
        existing = conn.execute(
            'SELECT 1 FROM kostenstellen_archive WHERE archiv_jahr = ? LIMIT 1',
            (jahr,),
        ).fetchone()
        if existing:
            return respond_error('Für dieses Jahr besteht bereits ein Archiv.')

        # akt. KSt. archivieren (immer)
        conn.execute(
            """
            INSERT INTO kostenstellen_archive (archiv_jahr, kostenstellen_nummer, kostenstellen_bezeichnung)
            SELECT ?, kostenstellen_nummer, kostenstellen_bezeichnung FROM kostenstellen
            """,
            (jahr,),
        )

        if jahr == active_year:
            # nur fürs aktive Jahr: KSt. ersetzen
            conn.execute('DELETE FROM kostenstellen')
            conn.executemany(
                'INSERT INTO kostenstellen (kostenstellen_nummer, kostenstellen_bezeichnung) VALUES (?, ?)',
                parsed_entries,
            )

        conn.commit()

    # nur wenn aktives Jahr abgeschlossen wird → auf nächstes Jahr umschalten
    if jahr == active_year:
        try:
            set_active_inventur_year(jahr + 1)
        except sqlite3.Error:
            app.logger.exception("Konnte active_inventur_year nicht setzen")

    audit_log('jahresabschluss', f'Jahr: {jahr}')
    return jsonify({'success': True, 'message': 'Jahresabschluss erfolgreich durchgeführt'})


@app.route('/update_article', methods=['POST'])
@require_admin
@require_csrf
def update_article():

    data = request.get_json(silent=True) or {}
    original_artikelname = (data.get('originalArticleName') or '').strip()
    new_artikelname = (data.get('newArticleName') or '').strip()
    new_einheit = (data.get('newUnit') or '').strip()
    new_price_raw = data.get('newPrice')

    if not (original_artikelname and new_artikelname and new_einheit):
        return respond_error('Ungültige Eingabedaten')

    try:
        new_price = parse_decimal(new_price_raw, 'Preis', allow_zero=True)
    except ValueError as exc:
        return respond_error(str(exc))

    with db_connection(ARTICLE_DB) as conn:
        conn.execute(
            'UPDATE artikel SET artikelname = ?, einheit = ?, preis = ? WHERE artikelname = ? COLLATE NOCASE',
            (
                new_artikelname,
                new_einheit,
                float(quantize_currency(new_price)),
                original_artikelname,
            ),
        )
        conn.commit()

    audit_log('artikel_aktualisiert', f'Alt: {original_artikelname}, Neu: {new_artikelname}')
    return jsonify({'success': True, 'message': 'Artikel erfolgreich aktualisiert'})


@app.route('/delete_article', methods=['POST'])
@require_admin
@require_csrf
def delete_article():
    artikelname = (request.get_json(silent=True) or {}).get('articleName', '').strip()

    if not artikelname:
        return respond_error('Ungültige Eingabedaten')

    with db_connection(ARTICLE_DB) as conn:
        conn.execute('DELETE FROM artikel WHERE artikelname = ? COLLATE NOCASE', (artikelname,))
        conn.commit()

    audit_log('artikel_geloescht', f'Artikel: {artikelname}')
    return jsonify({'success': True, 'message': 'Artikel erfolgreich geloescht'})


@app.route('/add_article', methods=['POST'])
@require_admin
@require_csrf
def add_article():
    data = request.get_json(silent=True) or {}
    artikelname = (data.get('name') or '').strip()
    einheit = (data.get('einheit') or '').strip()
    preis_raw = data.get('preis')

    if not artikelname or not einheit:
        return respond_error('Ungültige Eingabedaten')

    try:
        preis = parse_decimal(preis_raw, 'Preis')
    except ValueError as exc:
        return respond_error(str(exc))

    with db_connection(ARTICLE_DB) as conn:
        exists = conn.execute(
            'SELECT 1 FROM artikel WHERE artikelname = ? COLLATE NOCASE', (artikelname,)
        ).fetchone()
        if exists:
            return respond_error('Artikelname bereits vorhanden')

        conn.execute(
            'INSERT INTO artikel (artikelname, einheit, preis) VALUES (?, ?, ?)',
            (artikelname, einheit, float(quantize_currency(preis))),
        )
        conn.commit()

    audit_log('artikel_hinzugefuegt', f'Artikel: {artikelname}, Einheit: {einheit}')
    return jsonify({'success': True, 'message': 'Artikel erfolgreich hinzugefuegt'})


@app.route('/uebersicht_inventur')
def uebersicht_inventur():
    try:
        # Jahr-Filter aus Query (z. B. /uebersicht_inventur?jahr=2024). "all" oder leer = alle Jahre.
        jahr_param = (request.args.get('jahr') or '').strip()
        selected_year = None
        if jahr_param and jahr_param.lower() != 'all':
            try:
                selected_year = int(jahr_param)
            except ValueError:
                selected_year = None  # ungültig -> behandle wie "alle"

        with db_connection(COST_CENTER_DB) as conn:
            if selected_year is not None:
                # Nur dieses Jahr laden
                rows = conn.execute(
                    "SELECT rowid as _rid, *  FROM inventur WHERE jahr = ? ORDER BY kostenstelle, jahr, artikelname",
                    (selected_year,),
                ).fetchall()
            else:
                # Alle Jahre (default, rückwärtskompatibel)
                rows = conn.execute(
                    "SELECT rowid as _rid, * FROM inventur ORDER BY kostenstelle, jahr, artikelname"
                ).fetchall()

        gruppierte_inventur = group_inventur_rows(rows)
        alle_kostenstellen = sorted(gruppierte_inventur.keys())
        verfuegbare_jahre = fetch_available_years()

        archived_years = fetch_archived_years()
        archived_years_str = {str(y) for y in archived_years}

        return render_template(
            'uebersicht_inventur.html',
            gruppierte_inventur=gruppierte_inventur,
            alle_kostenstellen=alle_kostenstellen,
            verfuegbare_jahre=verfuegbare_jahre,
            ausgewaehltes_jahr=selected_year,
            active_inventur_year=get_active_inventur_year(),
            archived_years_str=archived_years_str,   # <--- NEU
        )
    except Exception:
        app.logger.exception("Fehler in uebersicht_inventur")
        return Response("Internal Server Error", status=500)




@app.route('/export_inventur_selection', methods=['POST'])
@require_csrf
def export_inventur_selection():
    # Eingaben
    jahr_param = (request.form.get('jahr') or '').strip()
    selected_year = None
    if jahr_param and jahr_param.lower() != 'all':
        try:
            selected_year = int(jahr_param)
        except ValueError:
            selected_year = None

    # kostenstellen: 'all' oder Liste
    # akzeptiere sowohl kostenstellen als auch kostenstellen[]
    raw_all = request.form.get('kostenstellen')
    raw_list = request.form.getlist('kostenstellen[]')
    if raw_all and raw_all.lower() == 'all':
        kst_filter = None  # keine Einschränkung
    else:
        kst_filter = [s.strip() for s in raw_list if str(s).strip()]

    # Daten laden
    base_sql = ('SELECT kostenstelle, jahr, artikelname, einheit, menge, preis, gesamtpreis '
                'FROM inventur')
    conds, params = [], []
    if selected_year is not None:
        conds.append('jahr = ?')
        params.append(selected_year)
    if kst_filter:
        placeholders = ','.join(['?'] * len(kst_filter))
        conds.append(f'kostenstelle IN ({placeholders})')
        params.extend(kst_filter)
    where = (' WHERE ' + ' AND '.join(conds)) if conds else ''
    order = ' ORDER BY kostenstelle, jahr, artikelname'

    with db_connection(COST_CENTER_DB) as conn:
        rows = conn.execute(base_sql + where + order, tuple(params)).fetchall()

    # Gruppieren & Excel bauen (gleich wie /export_inventur)
    gruppierte = group_inventur_rows(rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Inventurdaten'
    bold_font = Font(bold=True)
    header_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    currency_format = 'EUR #,##0.00'

    gesamtsumme_aller = Decimal('0')

    for kostenstelle, details in gruppierte.items():
        ws.append([kostenstelle])
        ws.cell(row=ws.max_row, column=1).font = bold_font

        ws.append(['Jahr', 'Artikelname', 'Menge', 'Einheit', 'Preis', 'Gesamtpreis'])
        for cell in ws[ws.max_row]:
            cell.font = bold_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        for material in details['materialien']:
            jahr = material['jahr'] if 'jahr' in material.keys() else material[2]
            artikelname = material['artikelname'] if 'artikelname' in material.keys() else material[3]
            menge = material['menge'] if 'menge' in material.keys() else material[5]
            einheit = material['einheit'] if 'einheit' in material.keys() else material[4]
            preis = material['preis'] if 'preis' in material.keys() else material[6]
            gesamtpreis = material['gesamtpreis'] if 'gesamtpreis' in material.keys() else material[7]
            ws.append([jahr, artikelname, menge, einheit, preis, gesamtpreis])

        gesamt = quantize_currency(details['gesamtsumme'])
        ws.append(['', '', '', '', 'Gesamt:', float(gesamt)])
        ws.cell(row=ws.max_row, column=5).font = bold_font
        ws.cell(row=ws.max_row, column=6).font = bold_font

        gesamtsumme_aller += gesamt

    ws.append(['', '', '', '', 'Gesamtsumme aller Kostenstellen:', float(quantize_currency(gesamtsumme_aller))])
    ws.cell(row=ws.max_row, column=5).font = bold_font
    ws.cell(row=ws.max_row, column=6).font = bold_font

    # optional: Zusammenfassungssheet
    ws_summary = wb.create_sheet(title='Uebersicht Gesamtsummen')
    ws_summary.append(['Kostenstelle', 'Gesamtsumme'])
    for cell in ws_summary[1]:
        cell.font = bold_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for kostenstelle, details in gruppierte.items():
        ws_summary.append([kostenstelle, float(quantize_currency(details['gesamtsumme']))])

    ws_summary.append(['Gesamtsumme aller Kostenstellen:', float(quantize_currency(gesamtsumme_aller))])
    for cell in ws_summary[ws_summary.max_row]:
        cell.font = bold_font
        cell.fill = header_fill

    def format_worksheet(sheet, currency_columns):
        for col_index in range(1, sheet.max_column + 1):
            column_letter = get_column_letter(col_index)
            max_length = 0
            for cell in sheet[column_letter]:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
                if column_letter in currency_columns:
                    cell.number_format = currency_format
            sheet.column_dimensions[column_letter].width = min(max_length + 2, 40)

    format_worksheet(ws, {'E', 'F'})
    format_worksheet(ws_summary, {'B'})

    # Dateiname
    fname = 'Inventur_Auswahl'
    if selected_year is not None:
        fname += f'_{selected_year}'
    if kst_filter and len(kst_filter) == 1:
        # ein bisschen säubern
        safe = re.sub(r'[^0-9A-Za-z_\-]+', '_', kst_filter[0])
        fname += f'_{safe}'
    fname += '.xlsx'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=fname,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )



@app.route('/export_inventur', methods=['POST'])
@require_csrf
def export_inventur():
    """Vollständiger Export aller Inventurdaten (alle Kostenstellen)."""
    jahr_param = (request.form.get('jahr') or '').strip()
    selected_year = None
    if jahr_param and jahr_param.lower() != 'all':
        try:
            selected_year = int(jahr_param)
        except ValueError:
            selected_year = None
    return _export_inventur_core(selected_year=selected_year, kst_filter=None)


def _export_inventur_core(selected_year, kst_filter):
    """Erstellt den Excel-Export für die angegebenen Filter."""
    base_sql = ('SELECT kostenstelle, jahr, artikelname, einheit, menge, preis, gesamtpreis '
                'FROM inventur')
    conds, params = [], []
    if selected_year is not None:
        conds.append('jahr = ?')
        params.append(selected_year)
    if kst_filter:
        placeholders = ','.join(['?'] * len(kst_filter))
        conds.append(f'kostenstelle IN ({placeholders})')
        params.extend(kst_filter)
    where = (' WHERE ' + ' AND '.join(conds)) if conds else ''
    order = ' ORDER BY kostenstelle, jahr, artikelname'

    with db_connection(COST_CENTER_DB) as conn:
        rows = conn.execute(base_sql + where + order, tuple(params)).fetchall()

    gruppierte = group_inventur_rows(rows)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Inventurdaten'
    bold_font = Font(bold=True)
    header_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    currency_format = 'EUR #,##0.00'
    gesamtsumme_aller = Decimal('0')

    for kostenstelle, details in gruppierte.items():
        ws.append([kostenstelle])
        ws.cell(row=ws.max_row, column=1).font = bold_font
        ws.append(['Jahr', 'Artikelname', 'Menge', 'Einheit', 'Preis', 'Gesamtpreis'])
        for cell in ws[ws.max_row]:
            cell.font = bold_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
        for material in details['materialien']:
            ws.append([
                material.get('jahr'), material.get('artikelname'), material.get('menge'),
                material.get('einheit'), material.get('preis'), material.get('gesamtpreis'),
            ])
        gesamt = quantize_currency(details['gesamtsumme'])
        ws.append(['', '', '', '', 'Gesamt:', float(gesamt)])
        ws.cell(row=ws.max_row, column=5).font = bold_font
        ws.cell(row=ws.max_row, column=6).font = bold_font
        gesamtsumme_aller += gesamt

    ws.append(['', '', '', '', 'Gesamtsumme aller Kostenstellen:', float(quantize_currency(gesamtsumme_aller))])
    for cell in ws[ws.max_row]:
        cell.font = bold_font

    for col_index in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_index)
        max_len = max((len(str(c.value)) for c in ws[col_letter] if c.value is not None), default=0)
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)
        if col_letter in {'E', 'F'}:
            for cell in ws[col_letter]:
                cell.number_format = currency_format

    fname = 'Inventur_Komplett'
    if selected_year is not None:
        fname += f'_{selected_year}'
    fname += '.xlsx'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=fname,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def fetch_available_years():
    with db_connection(COST_CENTER_DB) as conn:
        rows = conn.execute("SELECT DISTINCT jahr FROM inventur ORDER BY jahr DESC").fetchall()
    years = []
    for r in rows:
        try:
            years.append(int(r["jahr"] if hasattr(r, "keys") else r[0]))
        except Exception:
            continue
    return years

def fetch_archived_years():
    with db_connection(COST_CENTER_DB) as conn:
        rows = conn.execute("SELECT DISTINCT archiv_jahr FROM kostenstellen_archive").fetchall()
    vals = []
    for r in rows:
        try:
            vals.append(int(r["archiv_jahr"] if hasattr(r, "keys") else r[0]))
        except Exception:
            continue
    return sorted(set(vals), reverse=True)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000)
